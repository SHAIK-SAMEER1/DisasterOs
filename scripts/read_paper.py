import os
import docx

base_dir = r"C:\Users\SAMEER\.gemini\antigravity\scratch\disasteros"
path = os.path.join(base_dir, "24BDS1154_24BDS1171_Project1_Paper1-1-1.docx")
if os.path.exists(path):
    doc = docx.Document(path)
    print("Paragraphs count:", len(doc.paragraphs))
    print("Tables count:", len(doc.tables))
    for i, p in enumerate(doc.paragraphs[:40]):
        if p.text.strip():
            print(f"[{i}] {p.text.strip()[:140]}")
else:
    print("File not found:", path)
